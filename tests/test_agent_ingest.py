"""
test_agent_ingest.py
====================
Loading a new period from uploaded CSVs.

THE PROPERTY THIS SUITE PROVES
-------------------------------
An upload either updates **everything** or it is refused, and the refusal names
the row and the reason.

Both halves matter. A partial update -- a period that appears on one tab and
not another, or in the tables but not in the agent -- is worse than no upload
feature at all, because the tool looks like it worked. And a blind append is a
liability: if an upload can introduce an unknown department, a duplicate
period, or text where amounts belong, then every downstream guarantee is
resting on data nobody checked.

So the refusals are tested as carefully as the happy path, and the happy path
is followed all the way through to the agent producing an audited package for
the new month.
"""
import io
import os
import sys
import tempfile

import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "src")
sys.path.insert(0, SRC)

import run_pipeline as rp  # noqa: E402
from agent import tools as _tools  # noqa: E402,F401
from agent.ingest import (ESSENTIAL, SCHEMAS, classify, merge,  # noqa: E402
                          parse_upload, summarise)

NEW = "2026-01"


@pytest.fixture(scope="module")
def tables():
    return rp.load()


def csv_bytes(df) -> bytes:
    return df.to_csv(index=False).encode()


def month_slice(tables, name, source="2025-12", target=NEW):
    df = tables[name]
    out = df[df["month"].astype(str).str.startswith(source)].copy()
    out["month"] = target
    return out.astype(object)


@pytest.fixture
def good_files(tables):
    return [
        ("jan_actuals.csv", csv_bytes(month_slice(tables, "fact_actuals"))),
        ("jan_budget.csv", csv_bytes(month_slice(tables, "fact_budget"))),
        ("jan_headcount.csv", csv_bytes(month_slice(tables, "fact_headcount"))),
        ("jan_saas.csv", csv_bytes(month_slice(tables, "fact_saas_metrics"))),
        ("jan_saas_budget.csv",
         csv_bytes(month_slice(tables, "fact_saas_metrics_budget"))),
    ]


# --------------------------------------------------------------------------
# 1. classification
# --------------------------------------------------------------------------
def test_identical_schemas_are_disambiguated_by_filename():
    """actuals, budget and forecast share one schema EXACTLY, so the columns
    cannot tell them apart. Loading budget rows as actuals would silently
    corrupt every variance in the tool, so the filename decides."""
    cols = SCHEMAS["fact_actuals"]
    assert classify("jan_actuals.csv", cols)[0] == "fact_actuals"
    assert classify("2026-01_budget.csv", cols)[0] == "fact_budget"
    assert classify("forecast_2026.csv", cols)[0] == "fact_forecast"


def test_an_ambiguous_filename_is_refused_not_guessed():
    """Defaulting to actuals would be the worst possible failure: silent,
    total, and invisible until someone questioned a variance."""
    table, why = classify("Jan close v3 FINAL.csv", SCHEMAS["fact_actuals"])
    assert table is None
    assert "identical" in why and "filename" in why


def test_a_filename_naming_two_tables_is_refused():
    table, why = classify("actuals_and_budget.csv", SCHEMAS["fact_actuals"])
    assert table is None and "more than one" in why


def test_saas_files_also_need_a_hint_because_three_tables_share_the_schema():
    """SaaS metrics, its budget and its forecast are the same shape, so the
    same rule applies as for actuals/budget: a filename that says nothing is
    refused rather than defaulted."""
    table, why = classify("anything at all.csv", SCHEMAS["fact_saas_metrics"])
    assert table is None and "filename" in why

    # A hint of any strength resolves it.
    assert classify("arr.csv", SCHEMAS["fact_saas_metrics"])[0] == \
        "fact_saas_metrics"


def test_a_uniquely_shaped_table_needs_no_filename_hint():
    table, why = classify("whatever.csv", SCHEMAS["fact_headcount"])
    assert table == "fact_headcount" and "columns" in why


def test_the_most_specific_filename_match_wins():
    """"saas metrics budget" matches both the metrics and the budget table.
    The longer keyword is the more specific one."""
    cols = SCHEMAS["fact_saas_metrics"]
    assert classify("jan_saas_budget.csv", cols)[0] == "fact_saas_metrics_budget"
    assert classify("saas metrics forecast.csv", cols)[0] == \
        "fact_saas_metrics_forecast"
    assert classify("jan_saas.csv", cols)[0] == "fact_saas_metrics"


# --------------------------------------------------------------------------
# 2. refusals
# --------------------------------------------------------------------------
def _refuse(tables, df, name="jan_actuals.csv"):
    p = parse_upload([(name, csv_bytes(df))], tables)
    assert not p.acceptable, "this upload should have been refused"
    return " ".join(f.message for f in p.errors)


def test_a_period_that_already_exists_is_refused(tables):
    """An upload adds a period; it may not modify a closed one. Silently
    rewriting a closed month is how a reporting tool loses the trust the rest
    of its design earned."""
    df = month_slice(tables, "fact_actuals", target="2025-12")
    assert "already exist" in _refuse(tables, df)


def test_an_unknown_department_is_refused(tables):
    df = month_slice(tables, "fact_actuals")
    df.iloc[0, df.columns.get_loc("department_id")] = "LEGAL"
    msg = _refuse(tables, df)
    assert "Unknown department" in msg and "LEGAL" in msg
    assert "CORP" in msg, "the refusal should name the valid alternatives"


def test_an_unknown_account_is_refused(tables):
    df = month_slice(tables, "fact_actuals")
    df.iloc[0, df.columns.get_loc("account_id")] = "MADE_UP"
    assert "Unknown account" in _refuse(tables, df)


def test_a_non_numeric_amount_is_refused_with_the_row(tables):
    df = month_slice(tables, "fact_actuals")
    df.iloc[2, df.columns.get_loc("amount")] = "n/a"
    p = parse_upload([("jan_actuals.csv", csv_bytes(df))], tables)
    assert not p.acceptable
    err = next(f for f in p.errors)
    assert "not numbers" in err.message and err.rows, (
        "the refusal must name the offending row"
    )


def test_a_malformed_month_is_refused(tables):
    df = month_slice(tables, "fact_actuals")
    df["month"] = "Jan-2026"
    assert "YYYY-MM" in _refuse(tables, df)


def test_duplicate_rows_are_refused(tables):
    df = month_slice(tables, "fact_actuals")
    df = pd.concat([df, df.head(2)], ignore_index=True)
    assert "duplicate" in _refuse(tables, df)


def test_a_period_earlier_than_the_dataset_is_refused(tables):
    """Every month-over-month and trailing-window figure would silently span
    a gap."""
    df = month_slice(tables, "fact_actuals", target="2024-03")
    assert "already exist" in _refuse(tables, df) or "earlier" in _refuse(tables, df)


def test_a_missing_column_is_refused(tables):
    df = month_slice(tables, "fact_actuals").drop(columns=["amount"])
    assert "do not match any table" in _refuse(tables, df)


def test_two_files_of_the_same_table_are_refused(tables):
    df = month_slice(tables, "fact_actuals")
    p = parse_upload([("a_actuals.csv", csv_bytes(df)),
                      ("b_actuals.csv", csv_bytes(df))], tables)
    assert not p.acceptable
    assert any("both look like" in f.message for f in p.errors)


def test_non_utf8_is_refused_with_a_useful_instruction(tables):
    p = parse_upload([("jan_actuals.csv", b"\xff\xfe\x00bad")], tables)
    assert not p.acceptable
    assert "UTF-8" in " ".join(f.message for f in p.errors)


def test_no_files_is_refused(tables):
    assert not parse_upload([], tables).acceptable


# --------------------------------------------------------------------------
# 3. acceptance
# --------------------------------------------------------------------------
def test_accounting_formats_are_accepted(tables):
    """A finance export writes ($1,234) and $1,234 and means -1234 and 1234.
    Refusing those would fail on the most ordinary file a user could produce."""
    df = month_slice(tables, "fact_actuals")
    df["amount"] = df["amount"].map(
        lambda v: f"(${abs(float(v)):,.2f})" if float(v) < 0
        else f"${float(v):,.2f}")
    p = parse_upload([("jan_actuals.csv", csv_bytes(df))], tables)
    assert p.frames, " ".join(f.message for f in p.errors)
    parsed = p.frames["fact_actuals"]["amount"]
    original = month_slice(tables, "fact_actuals")["amount"].astype(float)
    assert parsed.round(2).tolist() == original.round(2).tolist()


def test_a_complete_upload_is_accepted(tables, good_files):
    p = parse_upload(good_files, tables)
    assert p.acceptable, " ".join(f.message for f in p.errors)
    assert p.new_periods == ["2026-01-01"]
    assert len(p.frames) == 5
    assert set(p.departments) == set(tables["dim_department"]["department_id"])
    assert "2026-01" in summarise(p)


def test_missing_optional_files_warn_with_the_consequence(tables):
    """"No SaaS metrics budget" means nothing to a reader. "The revenue
    volume-versus-price split will be unavailable" is actionable."""
    p = parse_upload(
        [("jan_actuals.csv", csv_bytes(month_slice(tables, "fact_actuals"))),
         ("jan_budget.csv", csv_bytes(month_slice(tables, "fact_budget")))],
        tables)
    assert p.acceptable
    warned = {f.table: f.message for f in p.warnings}
    assert "fact_headcount" in warned
    assert "unavailable" in warned["fact_headcount"]
    assert "refuse to produce a complete package" in \
        warned["fact_saas_metrics_budget"]


def test_essential_tables_are_named_when_absent(tables):
    p = parse_upload(
        [("jan_headcount.csv",
          csv_bytes(month_slice(tables, "fact_headcount")))], tables)
    warned = {f.table for f in p.warnings}
    for t in ESSENTIAL:
        assert t in warned


# --------------------------------------------------------------------------
# 4. the merge updates everything
# --------------------------------------------------------------------------
def test_merge_does_not_mutate_the_existing_tables(tables, good_files):
    before = len(tables["fact_actuals"])
    p = parse_upload(good_files, tables)
    merged = merge(tables, p)
    assert len(tables["fact_actuals"]) == before, "the original was mutated"
    assert len(merged["fact_actuals"]) > before


def test_the_date_dimension_gains_the_new_period(tables, good_files):
    """Without this every join against dim_date drops the new month and the
    period appears to have loaded while showing nothing."""
    merged = merge(tables, parse_upload(good_files, tables))
    months = set(merged["dim_date"]["month"].astype(str))
    assert "2026-01-01" in months
    row = merged["dim_date"][merged["dim_date"]["month"] == "2026-01-01"].iloc[0]
    assert row["fiscal_year"] == 2026 and row["quarter"] == "2026-Q1"
    assert row["month_number"] == 1


def test_the_new_period_reaches_the_computed_outputs(tables, good_files):
    merged = merge(tables, parse_upload(good_files, tables))
    out = rp.compute(merged)
    months = sorted(out["variance_detail"]["month"].astype(str).unique())
    assert months[-1] == "2026-01-01"
    assert not out["operating_metrics"].query("month == '2026-01-01'").empty
    assert not out["variance_by_department"].query(
        "month == '2026-01-01'").empty


def test_merge_is_refused_when_the_upload_was(tables):
    df = month_slice(tables, "fact_actuals", target="2025-12")
    p = parse_upload([("jan_actuals.csv", csv_bytes(df))], tables)
    with pytest.raises(ValueError, match="not accepted"):
        merge(tables, p)


@pytest.fixture
def restore_dataset():
    """Snapshot the committed CSVs and the built database, restore afterwards.

    The alternative was redirecting paths with monkeypatch, but the CSV and
    database locations are defined independently in THREE modules
    (agent.materialize, build_database, run_pipeline). Patching two of them
    produced an empty database and a test that failed for a reason unrelated
    to what it was testing.

    Snapshot-and-restore exercises the real path instead of a stitched-together
    one, which is the point: this test is about the production write, not about
    a redirected imitation of it.
    """
    import shutil

    from agent import materialize as mz

    syn, db = mz.SYN, mz.DB
    backup = tempfile.mkdtemp()
    shutil.copytree(syn, os.path.join(backup, "syn"))
    db_backup = None
    if os.path.exists(db):
        db_backup = os.path.join(backup, "fpa.duckdb")
        shutil.copy2(db, db_backup)
    # The session database counts as part of the state under restoration.
    # Without this every test after an upload sees the uploaded periods,
    # because connect_readonly prefers a session when one exists -- which is
    # exactly how four unrelated suites started failing on row counts.
    mz.clear_session()
    try:
        yield
    finally:
        mz.clear_session()
        shutil.rmtree(syn, ignore_errors=True)
        shutil.copytree(os.path.join(backup, "syn"), syn)
        if db_backup:
            shutil.copy2(db_backup, db)
        elif os.path.exists(db):
            os.remove(db)
        shutil.rmtree(backup, ignore_errors=True)


def test_the_agent_produces_a_package_for_an_uploaded_period(
        tables, good_files, restore_dataset):
    """The end of the chain, and the point of the feature.

    A period that loads into the tables but not into the agent is the partial
    update this suite exists to prevent -- the tool would look like it worked
    while the agent silently knew nothing about the new month.
    """
    from agent import materialize as mz
    from agent.narrate import narrate
    from agent.orchestrator import Orchestrator
    from agent.packages import build_goal, variance_package_plan
    from guardrails import entity_audit as ea

    merged = merge(tables, parse_upload(good_files, tables))

    # The real write, against the real paths. The restore_dataset fixture puts
    # the committed CSVs and database back afterwards.
    mz.write_tables(merged, verbose=False)

    con = mz.connect_readonly()
    try:
        goal = build_goal(con, NEW)
        result = Orchestrator(con).run(variance_package_plan(goal), goal)
        assert result.complete, f"{result.outcome} {result.missing_sections}"

        cand = narrate(result, goal, None,
                       ea.canonical_entity_names(merged), mode="inject")
        assert cand.publishable and cand.matched
    finally:
        con.close()


# --------------------------------------------------------------------------
# 5. the app surfaces it
# --------------------------------------------------------------------------
def _app(monkeypatch):
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    for env in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
        monkeypatch.delenv(env, raising=False)
    sys.path.insert(0, os.path.join(HERE, "..", "eval"))
    at = AppTest.from_file(os.path.join(HERE, "..", "app", "app.py"),
                           default_timeout=300)
    at.run()
    return at


def test_the_sidebar_offers_an_upload(monkeypatch):
    at = _app(monkeypatch)
    assert not at.exception
    assert at.get("file_uploader"), "no upload control"


def test_an_uploaded_period_extends_the_reporting_month_list(
        tables, good_files, monkeypatch):
    """The property the whole feature rests on: an upload updates EVERYTHING.

    A period that loaded into the tables but not the month selector would be
    invisible, and one that reached the selector but not the aggregates would
    be worse -- the tool would look like it worked.
    """
    at = _app(monkeypatch)
    before = list(at.sidebar.selectbox[0].options)

    preview = parse_upload(good_files, tables)
    assert preview.acceptable
    at.session_state["uploaded_tables"] = merge(tables, preview)
    at.session_state["uploaded_periods"] = list(preview.new_periods)
    at = at.run()
    assert not at.exception

    after = list(at.sidebar.selectbox[0].options)
    assert len(after) == len(before) + 1
    assert after[-1] != before[-1], "the new period is not the latest"

    at = at.sidebar.selectbox[0].set_value("2026-01-01").run()
    assert not at.exception
    labels = {m.label for m in at.metric}
    assert "Revenue" in labels and "Operating income" in labels


def test_every_page_renders_for_an_uploaded_period(tables, good_files,
                                                   monkeypatch):
    """A period that appears on one tab and not another is the partial update
    this feature exists to avoid."""
    at = _app(monkeypatch)
    preview = parse_upload(good_files, tables)
    at.session_state["uploaded_tables"] = merge(tables, preview)
    at.session_state["uploaded_periods"] = list(preview.new_periods)
    at = at.run()
    at = at.sidebar.selectbox[0].set_value("2026-01-01").run()

    for page in ("Overview", "Variance", "Forecast", "Commentary"):
        page_at = at.sidebar.radio[0].set_value(page).run()
        assert not page_at.exception, f"{page}: {page_at.exception}"


def test_the_upload_states_the_synthetic_data_rule(monkeypatch):
    """An upload control on a public link, operated by someone with a
    clearance, needs the rule stated at the control itself."""
    import re

    at = _app(monkeypatch)
    text = " ".join(re.sub("<[^>]+>", "", m.value) for m in at.markdown)
    assert "Synthetic or public data only" in text
    assert "never anything from an employer" in text.lower()


# --------------------------------------------------------------------------
# 6. the template workbook
# --------------------------------------------------------------------------
def _fill_template(tables, raw: bytes, target=NEW, source="2025-12",
                   sheets=("fact_actuals", "fact_budget", "fact_headcount",
                           "fact_saas_metrics", "fact_saas_metrics_budget")):
    """Fill the template the way a user would: replace the example row."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    for name in sheets:
        ws = wb[name]
        headers = [c.value for c in ws[4]
                   if c.value and c.value != "delete_this_example_row"]
        ws.delete_rows(5, ws.max_row)
        src = tables[name]
        src = src[src["month"].astype(str).str.startswith(source)]
        for i, (_, row) in enumerate(src.iterrows(), start=5):
            for j, col in enumerate(headers, start=1):
                if col == "month":
                    value = target
                elif col in ("department_id", "account_id"):
                    value = str(row[col])
                else:
                    value = float(row[col])
                ws.cell(row=i, column=j, value=value)
    for extra in set(wb.sheetnames) - set(sheets) - {"Instructions",
                                                     "Valid codes"}:
        wb.remove(wb[extra])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def template(tables):
    from agent.ingest import build_template

    return build_template(tables)


def test_the_template_has_a_sheet_per_table(template):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(template))
    assert "Instructions" in wb.sheetnames
    assert "Valid codes" in wb.sheetnames
    for required in ("fact_actuals", "fact_budget"):
        assert required in wb.sheetnames


def test_each_template_sheet_carries_its_headers_and_an_example(template):
    """A reader copies a row far more reliably than they interpret a schema."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(template))
    for name in ("fact_actuals", "fact_headcount", "fact_saas_metrics"):
        ws = wb[name]
        headers = [c.value for c in ws[4] if c.value]
        for col in SCHEMAS[name]:
            assert col in headers, f"{name} missing header {col}"
        example = [c.value for c in ws[5] if c.value is not None]
        assert example, f"{name} has no example row"


def test_the_template_lists_the_valid_codes(tables, template):
    """"Unknown department" is only a helpful refusal if the valid set is
    somewhere the user can see it."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(template))["Valid codes"]
    listed = {ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)}
    for dept in tables["dim_department"]["department_id"]:
        assert str(dept) in listed


def test_an_unfilled_template_is_refused_with_the_reason(tables, template):
    """The example row is tagged in a marker column outside every schema.
    Without it an untouched template uploaded cleanly and created a period
    containing one invented row -- the example silently becoming data."""
    p = parse_upload([("load-a-period-template.xlsx", template)], tables)
    assert not p.acceptable
    msg = " ".join(f.message for f in p.errors)
    assert "not been filled in" in msg


def test_a_filled_template_round_trips(tables, template):
    """The whole point: download, fill, upload, and the period loads."""
    filled = _fill_template(tables, template)
    p = parse_upload([("January 2026 close.xlsx", filled)], tables)
    assert p.acceptable, " ".join(f.message for f in p.errors)
    assert p.new_periods == ["2026-01-01"]
    assert set(p.frames) == {"fact_actuals", "fact_budget", "fact_headcount",
                             "fact_saas_metrics", "fact_saas_metrics_budget"}

    out = rp.compute(merge(tables, p))
    assert sorted(out["variance_detail"]["month"].astype(str))[-1] == \
        "2026-01-01"


def test_sheet_names_resolve_what_filenames_cannot(tables, template):
    """actuals and budget share one schema, so a loose CSV called "close
    final.csv" cannot be placed. A named sheet always can -- which is why the
    workbook is the recommended path."""
    filled = _fill_template(tables, template)
    p = parse_upload([("close final v3.xlsx", filled)], tables)
    assert p.acceptable
    assert "fact_actuals" in p.frames and "fact_budget" in p.frames

    loose = p.frames["fact_actuals"].to_csv(index=False).encode()
    bad = parse_upload([("close final v3.csv", loose)], tables)
    assert not bad.acceptable, "a loose CSV with no hint should be refused"


def test_a_workbook_with_no_recognised_sheets_is_refused(tables):
    import pandas as _pd

    buf = io.BytesIO()
    _pd.DataFrame({"a": [1]}).to_excel(buf, sheet_name="Sheet1", index=False)
    p = parse_upload([("unrelated.xlsx", buf.getvalue())], tables)
    assert not p.acceptable
    assert "No recognised sheets" in " ".join(f.message for f in p.errors)


def test_the_header_row_is_located_not_assumed(tables, template):
    """The template carries a purpose line and a REQUIRED marker above its
    headers. Reading row 1 as the header turns the description into a column
    name and the sheet is rejected for having the wrong columns."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(template))["fact_actuals"]
    assert ws.cell(row=1, column=1).value, "no purpose line above the headers"
    assert ws.cell(row=4, column=1).value == "month", \
        "headers are expected on row 4"

    filled = _fill_template(tables, template)
    assert parse_upload([("x.xlsx", filled)], tables).acceptable


def test_the_app_offers_the_template(monkeypatch):
    at = _app(monkeypatch)
    labels = [d.label for d in at.get("download_button")]
    assert any("template" in lbl.lower() for lbl in labels), labels


def test_the_uploader_accepts_workbooks_and_csvs(monkeypatch):
    at = _app(monkeypatch)
    up = at.get("file_uploader")
    assert up
    accepted = list(up[0].proto.type)
    assert ".xlsx" in accepted and ".csv" in accepted


# --------------------------------------------------------------------------
# 7. periods accumulate within a session
# --------------------------------------------------------------------------
def _month_workbook(tables, template, target, source):
    return _fill_template(tables, template, target=target, source=source)


def test_periods_accumulate_rather_than_replacing(tables, template):
    """A close cycle is a sequence, so loading January and then February is
    the ordinary case. The first version capped a session at one period for no
    reason other than simpler state.
    """
    running = tables
    loaded = []
    for target, source in (("2026-01", "2025-12"), ("2026-02", "2026-01"),
                           ("2026-03", "2026-02")):
        wb = _month_workbook(running, template, target, source)
        p = parse_upload([(f"{target}.xlsx", wb)], running)
        assert p.acceptable, " ".join(f.message for f in p.errors)
        running = merge(running, p)
        loaded += list(p.new_periods)

    months = sorted(rp.compute(running)["variance_detail"]["month"].astype(str)
                    .unique())
    assert months[-3:] == ["2026-01-01", "2026-02-01", "2026-03-01"]
    assert len(loaded) == 3


def test_a_second_upload_of_the_same_period_is_still_refused(tables, template):
    """Accumulating must not weaken the rule that a closed month is immutable."""
    first = _month_workbook(tables, template, "2026-01", "2025-12")
    running = merge(tables, parse_upload([("a.xlsx", first)], tables))

    again = parse_upload([("a.xlsx", first)], running)
    assert not again.acceptable
    assert "already exist" in " ".join(f.message for f in again.errors)


def test_a_gap_is_warned_about_not_refused(tables, template):
    """Loading January and then March leaves February missing, and every
    month-over-month and trailing-window figure then spans the hole. The data
    is not wrong, only incomplete, so it is surfaced loudly and left as the
    uploader's call."""
    jan = _month_workbook(tables, template, "2026-01", "2025-12")
    running = merge(tables, parse_upload([("jan.xlsx", jan)], tables))

    mar = _month_workbook(running, template, "2026-03", "2026-01")
    p = parse_upload([("mar.xlsx", mar)], running)

    assert p.acceptable, "a gap must not block a legitimate upload"
    gap = [f.message for f in p.warnings if "missing" in f.message]
    assert gap, "a gap must be surfaced"
    assert "2026-02" in gap[0]


def test_contiguous_uploads_raise_no_gap_warning(tables, template):
    jan = _month_workbook(tables, template, "2026-01", "2025-12")
    running = merge(tables, parse_upload([("jan.xlsx", jan)], tables))
    feb = _month_workbook(running, template, "2026-02", "2026-01")
    p = parse_upload([("feb.xlsx", feb)], running)
    assert not [f for f in p.warnings if "missing" in f.message]


def test_the_panel_keeps_offering_the_uploader_after_a_load(
        tables, template, monkeypatch):
    at = _app(monkeypatch)
    running = tables
    loaded = []
    for target, source in (("2026-01", "2025-12"), ("2026-02", "2026-01")):
        wb = _month_workbook(running, template, target, source)
        p = parse_upload([(f"{target}.xlsx", wb)], running)
        running = merge(running, p)
        loaded += list(p.new_periods)

    at.session_state["uploaded_tables"] = running
    at.session_state["uploaded_periods"] = sorted(set(loaded))
    at = at.run()
    assert not at.exception

    banner = " ".join(s.value for s in at.success)
    assert "January 2026" in banner and "February 2026" in banner
    assert at.get("file_uploader"), (
        "the uploader must remain available so the next month can be loaded"
    )


# --------------------------------------------------------------------------
# 8. an upload never touches the committed dataset
# --------------------------------------------------------------------------
def test_loading_a_period_does_not_modify_the_committed_csvs(
        tables, good_files):
    """The property the first implementation violated.

    write_tables originally wrote the merged tables back over data/synthetic,
    reasoning that the CSVs are the source of truth so writing them is the
    same path a fresh checkout takes. But an upload is SESSION state --
    provisional, discardable, scoped to one visitor -- and the CSVs are the
    committed dataset the whole project rests on.

    Loading a period silently rewrote the repository's source data. "Remove
    uploaded periods" cleared the session and could not put the CSVs back, and
    every later run saw a dataset nobody had committed: fact_actuals grew from
    624 rows to 650 and twenty-eight tests failed on the change, which is the
    only reason it surfaced.
    """
    import hashlib

    from agent import materialize as mz

    def fingerprint():
        h = hashlib.sha256()
        for name in sorted(os.listdir(mz.SYN)):
            if name.endswith(".csv"):
                with open(os.path.join(mz.SYN, name), "rb") as fh:
                    h.update(name.encode())
                    h.update(fh.read().replace(b"\r\n", b"\n"))
        return h.hexdigest()

    before = fingerprint()
    merged = merge(tables, parse_upload(good_files, tables))
    try:
        mz.write_tables(merged, verbose=False)
        assert fingerprint() == before, (
            "loading a period modified the committed CSVs"
        )
        con = mz.connect_readonly()
        try:
            latest = con.execute(
                "SELECT MAX(month) FROM out_variance_detail").fetchone()[0]
            assert str(latest).startswith("2026-01"), (
                "the agent should see the uploaded period"
            )
        finally:
            con.close()
    finally:
        mz.clear_session()


def test_discarding_an_upload_returns_the_agent_to_the_committed_data(
        tables, good_files):
    """Removing uploaded periods must clear the session database too, or the
    agent keeps serving periods every other page has just forgotten."""
    from agent import materialize as mz

    merged = merge(tables, parse_upload(good_files, tables))
    mz.write_tables(merged, verbose=False)
    assert mz.has_session()

    con = mz.connect_readonly()
    try:
        with_upload = con.execute(
            "SELECT MAX(month) FROM out_variance_detail").fetchone()[0]
    finally:
        con.close()
    assert str(with_upload).startswith("2026-01")

    assert mz.clear_session() is True
    assert not mz.has_session()

    con = mz.connect_readonly()
    try:
        restored = con.execute(
            "SELECT MAX(month) FROM out_variance_detail").fetchone()[0]
    finally:
        con.close()
    assert str(restored).startswith("2025-12"), (
        "discarding should return the agent to the committed dataset"
    )


def test_clearing_a_session_that_does_not_exist_is_harmless():
    from agent import materialize as mz

    mz.clear_session()
    assert mz.clear_session() is False


# --------------------------------------------------------------------------
# 9. departments read as names, not warehouse codes
# --------------------------------------------------------------------------
def test_department_tables_show_names_not_codes(monkeypatch):
    """CORP, RND and SM are the warehouse's identifiers, not names a reader
    outside the project would recognise.

    The agent's SQL tools already join dim_department for exactly this reason.
    The flagship pages read the computed frames directly and had no
    equivalent, so the same department appeared as "Research & Development" on
    one tab and "RND" on another.
    """
    at = _app(monkeypatch)
    at = at.sidebar.radio[0].set_value("Variance").run()
    assert not at.exception

    rollup = None
    drivers = None
    for element in at.dataframe:
        value = element.value
        if not hasattr(value, "columns"):
            continue
        cols = list(value.columns)
        if "Department" in cols:
            rollup = value
        if "Line item" in cols:
            drivers = value

    assert rollup is not None, "no department rollup table"
    names = set(rollup["Department"].astype(str))
    assert "Research & Development" in names
    assert "Sales & Marketing" in names
    assert not (names & {"RND", "SM", "CORP", "CS", "GA"}), (
        f"raw codes still shown: {names}"
    )

    assert drivers is not None, "no driver table"
    labels = " ".join(drivers["Line item"].astype(str))
    assert "(Corporate / Company)" in labels or \
        "(Sales & Marketing)" in labels, labels
    assert "(SM)" not in labels and "(RND)" not in labels


def test_dept_name_falls_back_to_the_code(monkeypatch):
    """An unrecognised department is still information, so it renders as its
    code rather than as a blank."""
    import importlib.util

    src = open(os.path.join(HERE, "..", "app", "app.py"),
               encoding="utf-8").read()
    start = src.index("def dept_name(")
    end = src.index("def money(")
    ns = {"DEPT_NAMES": {"SM": "Sales & Marketing"}}
    exec(src[start:end], ns)
    dept_name = ns["dept_name"]

    assert dept_name("SM") == "Sales & Marketing"
    assert dept_name("NEW_TEAM") == "NEW_TEAM"
    assert dept_name("NEW_TEAM", fallback=False) == ""
